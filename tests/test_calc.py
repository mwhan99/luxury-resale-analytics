"""The app's calculations must match the tracker spreadsheet's own formulas. Run: pytest"""
from pathlib import Path
import sys

import pytest
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import calc  # noqa: E402

DATA = Path(__file__).resolve().parents[1] / "data"
DEMO = DATA / "demo_tracker.xlsx"


def spreadsheet_summary(path):
    """Values the spreadsheet itself calculated (Summary tab, cached by Excel/LibreOffice)."""
    s = load_workbook(path, data_only=True)["Summary"]
    return {s.cell(r, 1).value: s.cell(r, 2).value for r in range(4, 13)}


def test_totals_match_spreadsheet_summary():
    k = calc.kpis(calc.load_tracker(DEMO))
    s = spreadsheet_summary(DEMO)
    assert k["bags"] == s["Bags tracked"]
    assert k["sold"] == s["Bags sold"]
    assert k["revenue"] == pytest.approx(s["Revenue from sold bags (£)"])
    assert k["profit"] == pytest.approx(s["Total profit (£)"])
    assert k["unsold_cost"] == pytest.approx(s["Unsold stock at cost (£)"])


def test_unsold_bags_have_no_profit():
    df = calc.load_tracker(DEMO)
    assert df.loc[~df["is_sold"], "profit"].isna().all()


def test_max_lot_price_hits_target_exactly():
    mx = calc.max_lot_price(13, 150, 0, 0.75, 160, target_roi=1.0)
    r = calc.lot_scenario(mx, 13, 150, 0, 0.75, 160)
    assert abs(r["roi"] - 1.0) < 1e-9


def test_buy_signal_needs_enough_comparables():
    assert calc.buy_signal(0.9, 1, 0.25)[0] == "CHECK"
    assert calc.buy_signal(0.3, 5, 0.25)[0] == "BUY"
    assert calc.buy_signal(0.1, 5, 0.25)[0] == "PASS"


def test_blank_template_loads_empty():
    assert len(calc.load_tracker(DATA / "blank_tracker.xlsx")) == 0
